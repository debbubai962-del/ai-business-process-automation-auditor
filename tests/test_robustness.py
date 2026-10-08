"""Hardening: the pipeline must degrade gracefully on unusual datasets, never crash or invent findings."""
import json
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from ai.brief import generate_brief
from ai.evidence import build_evidence
from ai.validate import validate_text
from ai.schema import flatten
from config.settings import Settings
from engines.mapping.detector import detect_mapping
from engines.mapping.event_log import build_event_log
from engines.quality.engine import clean_dataset
from reporting.audit import audit_dataframe
from reporting.excel_report import build_excel
from reporting.html_report import build_html_report
from reporting.safe import safe_cell, sanitize_frame
from services.bundle import assemble, compute_static

MESSY = Path(__file__).resolve().parents[1] / "sample_data" / "invoice_process_messy.csv"


@pytest.fixture(scope="module")
def raw():
    if not MESSY.exists():
        pytest.skip("Run sample_data/generate_demo_data.py first")
    return pd.read_csv(MESSY)


def full_pipeline(df: pd.DataFrame):
    q = clean_dataset(df)
    log = build_event_log(q.cleaned, {k: v for k, v in detect_mapping(q.cleaned).assignments.items() if v})
    bundle = assemble(compute_static(log), "test.csv", q, log)
    ev = build_evidence(bundle)
    brief = generate_brief(ev, Settings(openrouter_api_key=""))
    html = build_html_report(ev, brief, audit_dataframe(q), "now")
    xlsx = build_excel(ev, brief, audit_dataframe(q))
    assert brief.validation.ok, brief.validation.unsupported
    assert "<html" in html and len(xlsx) > 1000
    return bundle, ev, brief


def test_minimal_columns_only(raw):
    keep = ["invoice_id", "Activity", "Start time ", "End Time"]
    bundle, ev, brief = full_pipeline(raw[keep])
    assert ev["process_health"]["sla_compliance_percent"] is None     # no SLA column: reported as unavailable
    assert brief.brief.headline


def test_no_end_time_column(raw):
    bundle, ev, brief = full_pipeline(raw.drop(columns=["End Time"]))
    assert bundle.static.bottlenecks.approximate
    assert ev["caveats"]


def test_tiny_dataset(raw):
    ids = raw["invoice_id"].drop_duplicates().head(12)
    bundle, ev, brief = full_pipeline(raw[raw["invoice_id"].isin(ids)])
    assert ev["scope"]["cases"] <= 12


def test_single_department(raw):
    sub = raw[raw["Department"].astype(str).str.lower().str.contains("procure", na=False)]
    full_pipeline(sub)


def test_no_optional_dimensions(raw):
    full_pipeline(raw[["invoice_id", "Activity", "Start time ", "End Time", "SLA_hours"]])


def test_nothing_recognisable_is_rejected_cleanly():
    df = pd.DataFrame({"a": [1, 2, 3], "b": ["x", "y", "z"]})
    m = detect_mapping(df).assignments
    assert m["case_id"] is None and m["activity"] is None


def test_formula_injection_is_neutralised():
    assert safe_cell("=HYPERLINK(\"http://evil\")") == "'=HYPERLINK(\"http://evil\")"
    assert safe_cell("@SUM(A1)") == "'@SUM(A1)" and safe_cell("normal") == "normal" and safe_cell(5) == 5
    df = sanitize_frame(pd.DataFrame({"x": ["=1+1", "ok"], "n": [1, 2]}))
    assert df["x"].tolist() == ["'=1+1", "ok"]


def test_excel_export_has_no_live_formulas(raw):
    q = clean_dataset(raw)
    log = build_event_log(q.cleaned, {k: v for k, v in detect_mapping(q.cleaned).assignments.items() if v})
    ev = build_evidence(assemble(compute_static(log), "=cmd|' /C calc'!A0", q, log))
    brief = generate_brief(ev, Settings(openrouter_api_key=""))
    wb = load_workbook(BytesIO(build_excel(ev, brief, audit_dataframe(q))))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", f"{ws.title}!{c.coordinate} is a formula"