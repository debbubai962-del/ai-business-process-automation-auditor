"""Excel workbook of the findings. Aggregated results only."""
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from ai.brief import BriefResult
from reporting.safe import sanitize_frame


def _sheets(ev: dict, brief: BriefResult | None, audit: pd.DataFrame) -> dict[str, pd.DataFrame]:
    ph, dq, sim, sc = ev["process_health"], ev["data_quality"], ev["scenario_estimate"], ev["scope"]
    summary = [("Dataset", sc["dataset"]), ("Period", sc["period"]), ("Cases", sc["cases"]), ("Events", sc["events"])]
    summary += [(k.replace("_", " ").capitalize(), v) for k, v in ph.items()]
    summary += [(k.replace("_", " ").capitalize(), v) for k, v in dq.items() if not isinstance(v, list)]
    out = {
        "Summary": pd.DataFrame(summary, columns=["Measure", "Value"]),
        "Bottlenecks": pd.DataFrame(ev["bottlenecks"]),
        "Root causes": pd.DataFrame(ev["root_causes"]["at_primary_step"]["factors"]),
        "Opportunities": pd.DataFrame(ev["automation"]["recommended"]),
        "Not recommended": pd.DataFrame(ev["automation"]["not_recommended"]),
        "Scenario (estimate)": pd.DataFrame(
            [(k.replace("_", " ").capitalize(), ", ".join(v) if isinstance(v, list) else v)
             for k, v in sim.items() if k not in ("assumptions", "note")]
            + [("Assumption: " + k.replace("_", " "), v) for k, v in sim["assumptions"].items()],
            columns=["Item", "Value"]),
        "Audit trail": audit,
    }
    if brief is not None:
        b = brief.brief
        rows = [("Source", "AI" if brief.source == "ai" else "Rule-based (no AI)"), ("Headline", b.headline),
                ("Summary", b.executive_summary), ("Business impact (estimate)", b.business_impact)]
        rows += [("Issue", f"{i.title}. {i.detail}") for i in b.top_issues]
        rows += [("Hypothesis", f"{h.hypothesis} Evidence: {h.evidence}") for h in b.root_cause_hypotheses]
        rows += [("Action", f"{a.action}. {a.rationale}") for a in b.recommended_actions]
        rows += [("Plan: 30 days", x) for x in b.plan_30_days] + [("Plan: 60 days", x) for x in b.plan_60_days]
        rows += [("Plan: 90 days", x) for x in b.plan_90_days]
        out["Executive brief"] = pd.DataFrame(rows, columns=["Section", "Text"])
    return {k: sanitize_frame(v) for k, v in out.items() if len(v)}


def build_excel(ev: dict, brief: BriefResult | None, audit: pd.DataFrame) -> bytes:
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in _sheets(ev, brief, audit).items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0B1B3F")
            for col in ws.columns:
                width = min(70, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2)
                ws.column_dimensions[col[0].column_letter].width = max(12, width)
                for c in col[1:]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2"
    return buf.getvalue()